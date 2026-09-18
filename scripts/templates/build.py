"""
Builds the free templates sent by email from /free-templates:

  templates/<lang>/<open-to-buy file>.xlsx   open-to-buy planning model
  templates/<lang>/<checklist file>.xlsx     retail KPI dashboard checklist

    python scripts/templates/build.py

Texts per language live in TEXT below; the structure is the same in every
language. Yellow cells are inputs, white cells are formulas. The workbook is
flagged to recalculate on open, so Excel shows every result without a manual
recalculation. Formulas stay within Excel 2007 functions (IF, SUM, SUMIF,
COUNTIF, ROUND, ABS, OR, AND).
"""
from __future__ import annotations

import os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, "templates")

FONT = "Arial"
EMERALD = "0F5B4A"
INPUT_FILL = PatternFill("solid", fgColor="FFF4C2")
HEAD_FILL = PatternFill("solid", fgColor=EMERALD)
TOTAL_FILL = PatternFill("solid", fgColor="E8F1EE")
BLOCK_FILL = PatternFill("solid", fgColor="F6F4EF")
THIN = Side(style="thin", color="D5D2CA")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = "#,##0;-#,##0;-"
PCT = "0.0%"

# ---------------------------------------------------------------------------
# Texts
# ---------------------------------------------------------------------------

TEXT = {
    "en": {
        "otb_file": "open-to-buy-template.xlsx",
        "chk_file": "retail-kpi-dashboard-checklist.xlsx",
        "author": "Tatsiana Bandziuk · tatsianabandziuk.com",
        "sheets": {"guide": "Guide", "otb": "OTB", "checks": "Checks", "sizes": "Size & colour split", "defs": "Definitions"},
        "otb_title": "Open-to-buy plan by month, category and market",
        "otb_sub": "Yellow cells are inputs. White cells are formulas: do not type over them. All values at retail prices unless the line says cost or units.",
        "settings": ["Season", "Currency", "Current month (1–6)", "Drift tolerance", "Months"],
        "settings_vals": ["SS27", "EUR", 2, 0.10],
        "cur_note": "Months up to and including this number are closed: their actual sales and receipts must be filled in.",
        "months": ["Feb", "Mar", "Apr", "May", "Jun", "Jul"],
        "cols": ["Market", "Category", "Line", "Season total"],
        "lines": [
            "Planned sales",
            "Planned markdowns",
            "Planned closing stock",
            "Opening stock",
            "On order",
            "Open-to-buy",
            "Intake margin %",
            "Open-to-buy at cost",
            "Average unit cost",
            "Open-to-buy units",
            "Actual sales",
            "Actual receipts",
            "Check",
        ],
        "total": "All markets",
        "total_cat": "All categories",
        "markets": ["Market A", "Market B"],
        "cats": ["Dresses", "Knitwear"],
        "flag_over": "Over-bought",
        "flag_upd": "Update actuals",
        "flag_drift": "Sales off plan",
        "ok": "OK",
        "checks_title": "Three checks that stop the budget drifting mid-season",
        "checks_sub": "Counts of flagged months per block. Anything above zero needs a decision before the next order is placed.",
        "checks_cols": ["Market", "Category", "Over-bought months", "Actuals not updated", "Sales off plan", "Status"],
        "review": "Review",
        "checks_expl": [
            ("Over-bought", "Open-to-buy is below zero: stock on order and planned closing stock already exceed what the month needs. Cancel, move or re-phase orders."),
            ("Update actuals", "A closed month has no actual sales or receipts. The plan cannot be checked until both are filled in."),
            ("Sales off plan", "Actual sales differ from planned sales by more than the drift tolerance. Re-forecast the remaining months before buying more."),
        ],
        "sizes_title": "Size and colour split from cleaned sales history",
        "sizes_sub": "Use sales from weeks when every size was in stock, so a sold-out size does not look unpopular. The buy is split by the resulting shares.",
        "sizes_cat": "Category",
        "sizes_units": "Units to buy",
        "sizes_units_note": "Linked to the season open-to-buy units of the first block on the OTB sheet. Replace with another cell or a number.",
        "size_cols": ["Size", "Cleaned sales, units", "Share", "Units to buy"],
        "sizes": ["XS", "S", "M", "L", "XL"],
        "colour_cols": ["Colour", "Cleaned sales, units", "Share", "Units to buy"],
        "colours": ["Black", "Navy", "Beige", "Other"],
        "total_row": "Total",
        "rounding_note": "Rounded units can differ from the total by one or two; adjust the largest size.",
        "defs_title": "Definitions: how and by whom each figure is calculated",
        "defs_cols": ["Figure", "How it is calculated", "Owner"],
        "defs": [
            ("Planned sales", "Net sales at retail value planned for the month, after returns, before markdowns are deducted.", "Merchandise planner"),
            ("Planned markdowns", "Value lost to price reductions in the month: full price minus selling price, times units sold on reduction.", "Merchandise planner"),
            ("Planned closing stock", "Stock at retail value wanted at the end of the month, usually cover for the next few weeks of sales.", "Merchandise planner"),
            ("Opening stock", "First month: stock on hand at the start of the season. Later months: the planned closing stock of the month before.", "Formula"),
            ("On order", "Orders already placed for delivery in the month, at retail value.", "Buyer"),
            ("Open-to-buy", "Planned sales + planned markdowns + planned closing stock − opening stock − on order.", "Formula"),
            ("Intake margin %", "Planned margin on new buys: (retail price − cost price) ÷ retail price, both without VAT.", "Buyer and finance"),
            ("Open-to-buy at cost", "Open-to-buy × (1 − intake margin %): the budget the buyer can spend at cost.", "Formula"),
            ("Average unit cost", "Average cost price per unit of the category's new buys.", "Buyer"),
            ("Open-to-buy units", "Open-to-buy at cost ÷ average unit cost, rounded to whole units.", "Formula"),
            ("Actual sales, actual receipts", "Figures from the sales and stock reports for closed months, at retail value.", "Analyst"),
            ("Size and colour share", "Cleaned sales of a size or colour ÷ cleaned sales of the category.", "Merchandise planner"),
        ],
        "guide_title": "How to use the open-to-buy template",
        "guide": [
            ("Before the season", ""),
            ("1. Sales plan", "Enter planned sales by month for every market and category on the OTB sheet. Start from last season's sales for the same weeks."),
            ("2. Stock and markdown plan", "Enter planned markdowns and the closing stock you want each month. Enter the opening stock of the first month only; later months follow from the formula."),
            ("3. Budget", "Enter what is already on order, the intake margin and the average unit cost. The sheet returns the open-to-buy at retail, at cost and in units."),
            ("4. Orders", "Split the units by size and colour on the Size & colour split sheet, then place the orders and add them to On order."),
            ("During the season", ""),
            ("5. Monthly reconciliation", "When a month closes, fill in actual sales and receipts and raise the current month number. Read the Checks sheet before placing the next order."),
            ("6. Re-forecast", "If sales are off plan, change the remaining months' planned sales and markdowns first, then look at the open-to-buy again."),
            ("Handing the model over", ""),
            ("Owners", "Each figure has an owner on the Definitions sheet. Agree the names before the season starts."),
            ("Changes", "Adding a market or category: copy a full block of 13 rows below the last one and change the names. The All markets block picks it up automatically."),
            ("Legend", "Yellow cell: input. White cell: formula. Green header: section. Values in the example are illustrative; replace them with your own."),
        ],
        "chk_title": "Retail KPI dashboard checklist for Power BI",
        "chk_sub": "Twelve checks to run before a retail dashboard goes to the commercial team. Set a status for every line; anything not done is a reason to wait.",
        "chk_cols": ["#", "Area", "Check", "Why it matters", "Owner", "Status", "Notes"],
        "status": ["Done", "Not yet", "N/A"],
        "chk_summary": ["Done", "Applicable checks", "Progress", "Verdict"],
        "ready": "Ready to share",
        "not_ready": "Not ready yet",
        "chk": [
            ("KPI definitions", "Sell-through is defined and documented: formula, period, units or value.", "Two teams using two formulas will argue about the number instead of the decision."),
            ("KPI definitions", "Stock cover is shown per category and per market.", "A total cover figure hides the categories that are about to sell out or will end in markdown."),
            ("KPI definitions", "Full-price sales share and markdown depth sit next to each other.", "Sales growth bought with discounts looks like success until margin is checked."),
            ("KPI definitions", "Returns are in the same data model as sales.", "Net sales and sell-through are wrong for categories with high returns if returns live elsewhere."),
            ("KPI definitions", "Gross margin and GMROI use the same cost basis as finance.", "Different cost prices make the dashboard and the monthly finance report disagree."),
            ("Data model", "One date table marked as a date table, with weeks matching the trading calendar.", "Week-on-week and year-on-year comparisons shift when calendars differ."),
            ("Data model", "Stock is loaded as snapshots and never summed across dates.", "Adding up stock over weeks gives numbers many times the real inventory."),
            ("Data model", "One reporting currency, with the exchange rates and their date written down.", "Markets cannot be compared or added up in different currencies."),
            ("Access", "Row-level security per market is set and tested with View as role.", "Markets must see only their own figures, and a wrong role leaks data."),
            ("Refresh", "Scheduled refresh works and the last refresh date is shown on the report.", "People stop trusting a report once they find out it showed last week's data."),
            ("Ownership", "Every report page has one named owner.", "Questions and errors need one person to go to, not a shared inbox."),
            ("Reconciliation", "Totals for one closed week match the finance or ERP report.", "One reconciled week is the fastest proof that the model is right."),
        ],
    },
    "pl": {
        "otb_file": "szablon-open-to-buy.xlsx",
        "chk_file": "checklista-dashboardu-kpi.xlsx",
        "author": "Tatsiana Bandziuk · tatsianabandziuk.com",
        "sheets": {"guide": "Instrukcja", "otb": "OTB", "checks": "Kontrole", "sizes": "Rozmiary i kolory", "defs": "Definicje"},
        "otb_title": "Budżet open-to-buy według miesięcy, kategorii i rynków",
        "otb_sub": "Żółte komórki to dane wejściowe. Białe to formuły: nie nadpisuj ich. Wartości w cenach detalicznych, chyba że wiersz mówi o koszcie lub sztukach.",
        "settings": ["Sezon", "Waluta", "Bieżący miesiąc (1–6)", "Tolerancja odchylenia", "Miesiące"],
        "settings_vals": ["WL27", "PLN", 2, 0.10],
        "cur_note": "Miesiące do tego numeru włącznie są zamknięte: trzeba w nich uzupełnić rzeczywistą sprzedaż i dostawy.",
        "months": ["lut", "mar", "kwi", "maj", "cze", "lip"],
        "cols": ["Rynek", "Kategoria", "Pozycja", "Suma sezonu"],
        "lines": [
            "Plan sprzedaży",
            "Plan obniżek",
            "Plan zapasu końcowego",
            "Zapas początkowy",
            "Zamówione",
            "Open-to-buy",
            "Marża wejściowa %",
            "Open-to-buy w koszcie",
            "Średni koszt sztuki",
            "Open-to-buy w sztukach",
            "Sprzedaż rzeczywista",
            "Dostawy rzeczywiste",
            "Kontrola",
        ],
        "total": "Wszystkie rynki",
        "total_cat": "Wszystkie kategorie",
        "markets": ["Rynek A", "Rynek B"],
        "cats": ["Sukienki", "Dzianina"],
        "flag_over": "Przekroczony budżet",
        "flag_upd": "Uzupełnij dane",
        "flag_drift": "Sprzedaż poza planem",
        "ok": "OK",
        "checks_title": "Trzy kontrole, które zatrzymują rozjazd budżetu w sezonie",
        "checks_sub": "Liczba oznaczonych miesięcy w każdym bloku. Wynik powyżej zera wymaga decyzji przed kolejnym zamówieniem.",
        "checks_cols": ["Rynek", "Kategoria", "Miesiące z przekroczeniem", "Brak danych rzeczywistych", "Sprzedaż poza planem", "Status"],
        "review": "Do przeglądu",
        "checks_expl": [
            ("Przekroczony budżet", "Open-to-buy jest poniżej zera: zamówienia i plan zapasu końcowego przekraczają potrzeby miesiąca. Anuluj, przesuń lub rozłóż zamówienia."),
            ("Uzupełnij dane", "Zamknięty miesiąc nie ma rzeczywistej sprzedaży lub dostaw. Planu nie da się sprawdzić, dopóki obie wartości nie są wpisane."),
            ("Sprzedaż poza planem", "Sprzedaż rzeczywista różni się od planu bardziej niż tolerancja. Zaktualizuj prognozę na kolejne miesiące przed dalszymi zakupami."),
        ],
        "sizes_title": "Podział na rozmiary i kolory według oczyszczonej sprzedaży",
        "sizes_sub": "Bierz sprzedaż z tygodni, w których wszystkie rozmiary były dostępne, żeby wyprzedany rozmiar nie wyglądał na niepopularny. Zakup dzieli się według otrzymanych udziałów.",
        "sizes_cat": "Kategoria",
        "sizes_units": "Sztuki do zakupu",
        "sizes_units_note": "Połączone z sumą open-to-buy w sztukach pierwszego bloku w arkuszu OTB. Można podać inną komórkę lub liczbę.",
        "size_cols": ["Rozmiar", "Oczyszczona sprzedaż, szt.", "Udział", "Sztuki do zakupu"],
        "sizes": ["XS", "S", "M", "L", "XL"],
        "colour_cols": ["Kolor", "Oczyszczona sprzedaż, szt.", "Udział", "Sztuki do zakupu"],
        "colours": ["Czarny", "Granatowy", "Beżowy", "Inne"],
        "total_row": "Razem",
        "rounding_note": "Po zaokrągleniu suma może różnić się o jedną lub dwie sztuki; skoryguj największy rozmiar.",
        "defs_title": "Definicje: jak i przez kogo liczony jest każdy wskaźnik",
        "defs_cols": ["Wskaźnik", "Jak jest liczony", "Właściciel"],
        "defs": [
            ("Plan sprzedaży", "Planowana sprzedaż netto miesiąca w wartości detalicznej, po zwrotach, przed odjęciem obniżek.", "Planista"),
            ("Plan obniżek", "Wartość utracona na obniżkach w miesiącu: cena pełna minus cena sprzedaży, razy sztuki sprzedane po obniżce.", "Planista"),
            ("Plan zapasu końcowego", "Zapas w wartości detalicznej na koniec miesiąca, zwykle pokrycie sprzedaży kolejnych tygodni.", "Planista"),
            ("Zapas początkowy", "Pierwszy miesiąc: zapas na początek sezonu. Kolejne miesiące: plan zapasu końcowego z poprzedniego miesiąca.", "Formuła"),
            ("Zamówione", "Zamówienia już złożone z dostawą w danym miesiącu, w wartości detalicznej.", "Kupiec"),
            ("Open-to-buy", "Plan sprzedaży + plan obniżek + plan zapasu końcowego − zapas początkowy − zamówione.", "Formuła"),
            ("Marża wejściowa %", "Planowana marża na nowych zakupach: (cena detaliczna − cena zakupu) ÷ cena detaliczna, obie bez VAT.", "Kupiec i finanse"),
            ("Open-to-buy w koszcie", "Open-to-buy × (1 − marża wejściowa %): budżet, który kupiec może wydać w cenach zakupu.", "Formuła"),
            ("Średni koszt sztuki", "Średnia cena zakupu jednej sztuki nowych zakupów w kategorii.", "Kupiec"),
            ("Open-to-buy w sztukach", "Open-to-buy w koszcie ÷ średni koszt sztuki, zaokrąglone do pełnych sztuk.", "Formuła"),
            ("Sprzedaż i dostawy rzeczywiste", "Dane z raportów sprzedaży i zapasów za zamknięte miesiące, w wartości detalicznej.", "Analityk"),
            ("Udział rozmiaru i koloru", "Oczyszczona sprzedaż rozmiaru lub koloru ÷ oczyszczona sprzedaż kategorii.", "Planista"),
        ],
        "guide_title": "Jak korzystać z szablonu open-to-buy",
        "guide": [
            ("Przed sezonem", ""),
            ("1. Plan sprzedaży", "Wpisz planowaną sprzedaż według miesięcy dla każdego rynku i kategorii w arkuszu OTB. Zacznij od sprzedaży z tych samych tygodni poprzedniego sezonu."),
            ("2. Plan zapasu i obniżek", "Wpisz planowane obniżki i zapas, który chcesz mieć na koniec każdego miesiąca. Zapas początkowy wpisz tylko dla pierwszego miesiąca; kolejne wynikają z formuły."),
            ("3. Budżet", "Wpisz to, co jest już zamówione, marżę wejściową i średni koszt sztuki. Arkusz poda open-to-buy w cenach detalicznych, w koszcie i w sztukach."),
            ("4. Zamówienia", "Podziel sztuki na rozmiary i kolory w arkuszu Rozmiary i kolory, złóż zamówienia i dopisz je w wierszu Zamówione."),
            ("W trakcie sezonu", ""),
            ("5. Miesięczne uzgodnienie", "Po zamknięciu miesiąca wpisz sprzedaż i dostawy rzeczywiste i zwiększ numer bieżącego miesiąca. Przed kolejnym zamówieniem sprawdź arkusz Kontrole."),
            ("6. Nowa prognoza", "Jeśli sprzedaż odbiega od planu, najpierw zmień plan sprzedaży i obniżek na kolejne miesiące, a potem wróć do open-to-buy."),
            ("Przekazanie modelu", ""),
            ("Właściciele", "Każdy wskaźnik ma właściciela w arkuszu Definicje. Ustalcie osoby przed startem sezonu."),
            ("Zmiany", "Nowy rynek lub kategoria: skopiuj pełny blok 13 wierszy pod ostatnim i zmień nazwy. Blok Wszystkie rynki uwzględni go automatycznie."),
            ("Legenda", "Żółta komórka: dane wejściowe. Biała: formuła. Zielony nagłówek: sekcja. Wartości w przykładzie są poglądowe; zastąp je własnymi."),
        ],
        "chk_title": "Checklista dashboardu KPI dla handlu w Power BI",
        "chk_sub": "Dwanaście punktów do sprawdzenia, zanim dashboard trafi do zespołu handlowego. Ustaw status każdego punktu; każdy niespełniony to powód, żeby poczekać.",
        "chk_cols": ["#", "Obszar", "Punkt kontroli", "Dlaczego to ważne", "Właściciel", "Status", "Uwagi"],
        "status": ["Gotowe", "Jeszcze nie", "Nie dotyczy"],
        "chk_summary": ["Gotowe", "Punkty do sprawdzenia", "Postęp", "Werdykt"],
        "ready": "Można udostępnić",
        "not_ready": "Jeszcze nie gotowe",
        "chk": [
            ("Definicje KPI", "Sell-through jest zdefiniowany i opisany: formuła, okres, sztuki czy wartość.", "Dwa zespoły z dwiema formułami będą spierać się o liczbę, a nie o decyzję."),
            ("Definicje KPI", "Pokrycie zapasu jest pokazane według kategorii i rynku.", "Łączne pokrycie ukrywa kategorie, które zaraz się wyprzedadzą albo skończą na obniżkach."),
            ("Definicje KPI", "Udział sprzedaży w pełnej cenie i głębokość obniżek są obok siebie.", "Wzrost sprzedaży kupiony rabatami wygląda na sukces, dopóki nikt nie sprawdzi marży."),
            ("Definicje KPI", "Zwroty są w tym samym modelu danych co sprzedaż.", "Sprzedaż netto i sell-through są błędne w kategoriach z dużymi zwrotami, jeśli zwroty są gdzie indziej."),
            ("Definicje KPI", "Marża brutto i GMROI liczone są na tej samej bazie kosztowej co w finansach.", "Różne ceny zakupu sprawiają, że dashboard i miesięczny raport finansowy się nie zgadzają."),
            ("Model danych", "Jedna tabela dat oznaczona jako tabela dat, z tygodniami zgodnymi z kalendarzem handlowym.", "Porównania tydzień do tygodnia i rok do roku się przesuwają, gdy kalendarze się różnią."),
            ("Model danych", "Zapas jest ładowany jako stany na dzień i nigdy nie sumuje się go po datach.", "Suma zapasów z kilku tygodni daje wielokrotność rzeczywistego stanu."),
            ("Model danych", "Jedna waluta raportowa, z zapisanymi kursami i datą kursu.", "Rynków w różnych walutach nie da się porównać ani zsumować."),
            ("Dostęp", "Zabezpieczenia na poziomie wierszy (RLS) dla rynków są ustawione i przetestowane opcją Wyświetl jako rola.", "Rynki mają widzieć tylko swoje dane, a błędna rola ujawnia cudze."),
            ("Odświeżanie", "Zaplanowane odświeżanie działa, a data ostatniego odświeżenia jest widoczna w raporcie.", "Ludzie przestają ufać raportowi, gdy odkryją, że pokazywał dane sprzed tygodnia."),
            ("Właściciele", "Każda strona raportu ma jednego wskazanego z nazwiska właściciela.", "Pytania i błędy muszą trafiać do jednej osoby, a nie do wspólnej skrzynki."),
            ("Uzgodnienie", "Sumy za jeden zamknięty tydzień zgadzają się z raportem finansowym lub z ERP.", "Jeden uzgodniony tydzień to najszybszy dowód, że model liczy poprawnie."),
        ],
    },
    "ru": {
        "otb_file": "shablon-byudzheta-zakupok-open-to-buy.xlsx",
        "chk_file": "chek-list-dashborda-kpi.xlsx",
        "author": "Татьяна Бандюк · tatsianabandziuk.com",
        "sheets": {"guide": "Инструкция", "otb": "OTB", "checks": "Проверки", "sizes": "Размеры и цвета", "defs": "Определения"},
        "otb_title": "Бюджет закупок open-to-buy по месяцам, категориям и рынкам",
        "otb_sub": "Жёлтые ячейки — ввод данных. Белые — формулы, их не перезаписывайте. Все суммы в розничных ценах, если в строке не сказано о себестоимости или штуках.",
        "settings": ["Сезон", "Валюта", "Текущий месяц (1–6)", "Допустимое отклонение", "Месяцы"],
        "settings_vals": ["ВЛ27", "EUR", 2, 0.10],
        "cur_note": "Месяцы до этого номера включительно закрыты: в них нужно внести фактические продажи и поставки.",
        "months": ["фев", "мар", "апр", "май", "июн", "июл"],
        "cols": ["Рынок", "Категория", "Строка", "Итого за сезон"],
        "lines": [
            "План продаж",
            "План уценки",
            "План остатка на конец",
            "Остаток на начало",
            "Заказано",
            "Open-to-buy",
            "Маржинальность закупки %",
            "Open-to-buy по себестоимости",
            "Средняя себестоимость единицы",
            "Open-to-buy в штуках",
            "Продажи факт",
            "Поставки факт",
            "Проверка",
        ],
        "total": "Все рынки",
        "total_cat": "Все категории",
        "markets": ["Рынок A", "Рынок B"],
        "cats": ["Платья", "Трикотаж"],
        "flag_over": "Перезакуп",
        "flag_upd": "Внесите факт",
        "flag_drift": "Продажи вне плана",
        "ok": "OK",
        "checks_title": "Три проверки, которые не дают бюджету разойтись с фактом",
        "checks_sub": "Число отмеченных месяцев по каждому блоку. Всё, что больше нуля, требует решения до следующего заказа.",
        "checks_cols": ["Рынок", "Категория", "Месяцы с перезакупом", "Факт не внесён", "Продажи вне плана", "Статус"],
        "review": "Разобрать",
        "checks_expl": [
            ("Перезакуп", "Open-to-buy меньше нуля: заказы и плановый остаток уже больше, чем нужно месяцу. Отмените, перенесите или разнесите заказы по месяцам."),
            ("Внесите факт", "В закрытом месяце нет фактических продаж или поставок. План нельзя проверить, пока не внесены обе цифры."),
            ("Продажи вне плана", "Фактические продажи отличаются от плана больше, чем на допустимое отклонение. Сначала пересчитайте прогноз на оставшиеся месяцы, потом закупайте."),
        ],
        "sizes_title": "Размерные и цветовые доли по очищенным продажам",
        "sizes_sub": "Берите продажи за недели, когда в наличии были все размеры: так распроданный размер не выглядит непопулярным. Закупка делится по полученным долям.",
        "sizes_cat": "Категория",
        "sizes_units": "Штук к закупке",
        "sizes_units_note": "Связано с итогом open-to-buy в штуках первого блока на листе OTB. Можно заменить другой ячейкой или числом.",
        "size_cols": ["Размер", "Очищенные продажи, шт.", "Доля", "Штук к закупке"],
        "sizes": ["XS", "S", "M", "L", "XL"],
        "colour_cols": ["Цвет", "Очищенные продажи, шт.", "Доля", "Штук к закупке"],
        "colours": ["Чёрный", "Тёмно-синий", "Бежевый", "Другие"],
        "total_row": "Итого",
        "rounding_note": "После округления сумма может отличаться на одну-две штуки: поправьте самый большой размер.",
        "defs_title": "Определения: как и кем считается каждый показатель",
        "defs_cols": ["Показатель", "Как считается", "Кто отвечает"],
        "defs": [
            ("План продаж", "Плановые чистые продажи месяца в розничных ценах, после возвратов, до вычета уценки.", "Планировщик"),
            ("План уценки", "Потери от снижения цен за месяц: полная цена минус цена продажи, умноженные на штуки, проданные со скидкой.", "Планировщик"),
            ("План остатка на конец", "Остаток в розничных ценах, нужный на конец месяца, обычно запас на несколько следующих недель продаж.", "Планировщик"),
            ("Остаток на начало", "Первый месяц: остаток на начало сезона. Следующие месяцы: плановый остаток на конец предыдущего месяца.", "Формула"),
            ("Заказано", "Уже размещённые заказы с поставкой в этом месяце, в розничных ценах.", "Байер"),
            ("Open-to-buy", "План продаж + план уценки + план остатка на конец − остаток на начало − заказано.", "Формула"),
            ("Маржинальность закупки %", "Плановая маржинальность новых закупок: (розничная цена − закупочная цена) ÷ розничная цена, обе без НДС.", "Байер и финансы"),
            ("Open-to-buy по себестоимости", "Open-to-buy × (1 − маржинальность закупки %): бюджет, который байер может потратить в закупочных ценах.", "Формула"),
            ("Средняя себестоимость единицы", "Средняя закупочная цена единицы новых закупок категории.", "Байер"),
            ("Open-to-buy в штуках", "Open-to-buy по себестоимости ÷ средняя себестоимость единицы, округлённо до целых штук.", "Формула"),
            ("Продажи и поставки факт", "Данные отчётов по продажам и остаткам за закрытые месяцы, в розничных ценах.", "Аналитик"),
            ("Доля размера и цвета", "Очищенные продажи размера или цвета ÷ очищенные продажи категории.", "Планировщик"),
        ],
        "guide_title": "Как работать с шаблоном бюджета закупок",
        "guide": [
            ("До сезона", ""),
            ("1. План продаж", "Внесите план продаж по месяцам для каждого рынка и категории на листе OTB. Отправная точка — продажи тех же недель прошлого сезона."),
            ("2. План остатков и уценки", "Внесите план уценки и остаток, который нужен на конец каждого месяца. Остаток на начало вносится только для первого месяца, дальше его считает формула."),
            ("3. Бюджет", "Внесите уже размещённые заказы, маржинальность закупки и среднюю себестоимость единицы. Лист покажет open-to-buy в рознице, по себестоимости и в штуках."),
            ("4. Заказы", "Разложите штуки по размерам и цветам на листе «Размеры и цвета», разместите заказы и добавьте их в строку «Заказано»."),
            ("В течение сезона", ""),
            ("5. Сверка с фактом", "Когда месяц закрыт, внесите фактические продажи и поставки и увеличьте номер текущего месяца. Перед следующим заказом откройте лист «Проверки»."),
            ("6. Новый прогноз", "Если продажи ушли от плана, сначала измените план продаж и уценки на оставшиеся месяцы, потом снова смотрите на open-to-buy."),
            ("Передача модели команде", ""),
            ("Владельцы", "У каждого показателя есть владелец на листе «Определения». Согласуйте имена до начала сезона."),
            ("Изменения", "Новый рынок или категория: скопируйте полный блок из 13 строк под последним и поменяйте названия. Блок «Все рынки» учтёт его автоматически."),
            ("Обозначения", "Жёлтая ячейка — ввод. Белая — формула. Зелёный заголовок — раздел. Значения в примере условные, замените их своими."),
        ],
        "chk_title": "Чек-лист дашборда KPI для ритейла в Power BI",
        "chk_sub": "Двенадцать проверок перед тем, как дашборд увидит коммерческая команда. Поставьте статус каждой строке; любой невыполненный пункт — повод подождать.",
        "chk_cols": ["№", "Область", "Проверка", "Почему это важно", "Кто отвечает", "Статус", "Комментарий"],
        "status": ["Готово", "Ещё нет", "Не относится"],
        "chk_summary": ["Готово", "Применимых проверок", "Прогресс", "Итог"],
        "ready": "Можно показывать",
        "not_ready": "Пока рано",
        "chk": [
            ("Определения KPI", "Sell-through определён и описан: формула, период, штуки или деньги.", "Две команды с двумя формулами спорят о цифре, а не о решении."),
            ("Определения KPI", "Покрытие запасом показано по категориям и по рынкам.", "Общее покрытие скрывает категории, которые скоро закончатся или уйдут в уценку."),
            ("Определения KPI", "Доля продаж по полной цене и глубина уценки стоят рядом.", "Рост продаж за счёт скидок выглядит успехом, пока никто не посмотрел на маржу."),
            ("Определения KPI", "Возвраты находятся в той же модели данных, что и продажи.", "Чистые продажи и sell-through неверны для категорий с высокими возвратами, если возвраты лежат отдельно."),
            ("Определения KPI", "Валовая маржа и GMROI считаются от той же себестоимости, что и у финансов.", "Разная себестоимость — и дашборд расходится с ежемесячным финансовым отчётом."),
            ("Модель данных", "Одна таблица дат, отмеченная как таблица дат, с неделями по торговому календарю.", "Сравнения неделя к неделе и год к году сдвигаются, если календари разные."),
            ("Модель данных", "Остатки загружены как срезы на дату и никогда не суммируются по датам.", "Сумма остатков за несколько недель в разы больше реального запаса."),
            ("Модель данных", "Одна валюта отчёта, курсы и дата курса записаны.", "Рынки в разных валютах нельзя ни сравнить, ни сложить."),
            ("Доступ", "Безопасность на уровне строк (RLS) по рынкам настроена и проверена через «Просмотреть как роль».", "Каждый рынок должен видеть только свои цифры, а неправильная роль открывает чужие."),
            ("Обновление", "Плановое обновление работает, дата последнего обновления видна в отчёте.", "Отчёту перестают доверять, как только узнают, что он показывал данные недельной давности."),
            ("Владельцы", "У каждой страницы отчёта один владелец с именем.", "С вопросами и ошибками нужно идти к одному человеку, а не в общий ящик."),
            ("Сверка", "Итоги за одну закрытую неделю совпадают с финансовым отчётом или ERP.", "Одна сверенная неделя — самое быстрое доказательство, что модель считает верно."),
        ],
    },
}

# Illustrative values per block: sales, markdowns, closing stock, opening stock (month 1), on order,
# intake margin, unit cost, actual sales and receipts for the first two months.
EXAMPLE = [
    dict(sales=[42000, 48000, 55000, 60000, 52000, 38000], md=[1500, 1800, 2500, 4000, 9000, 12000],
         close=[120000, 130000, 125000, 110000, 80000, 50000], open=110000, onorder=[30000, 40000, 35000, 20000, 10000, 0],
         margin=0.62, cost=24, actual=[40500, 51000], receipts=[29000, 41500]),
    dict(sales=[26000, 24000, 18000, 12000, 9000, 7000], md=[1000, 1500, 2500, 3500, 3000, 2500],
         close=[70000, 60000, 45000, 30000, 20000, 12000], open=75000, onorder=[15000, 8000, 5000, 0, 0, 0],
         margin=0.60, cost=19, actual=[27200, 21000], receipts=[15000, 8000]),
    dict(sales=[18000, 21000, 24000, 26000, 22000, 16000], md=[600, 800, 1100, 1800, 4000, 5500],
         close=[52000, 56000, 54000, 47000, 35000, 22000], open=48000, onorder=[14000, 18000, 15000, 9000, 4000, 0],
         margin=0.62, cost=24, actual=[16100, 19800], receipts=[14000, None]),
    dict(sales=[11000, 10000, 7500, 5000, 3800, 3000], md=[400, 650, 1100, 1500, 1300, 1100],
         close=[30000, 26000, 19000, 13000, 8500, 5000], open=32000, onorder=[6000, 3500, 2000, 0, 0, 0],
         margin=0.60, cost=19, actual=[11300, 10400], receipts=[6000, 3500]),
]
SIZE_SALES = [90, 260, 340, 220, 90]
COLOUR_SALES = [420, 260, 190, 130]

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def font(**kw):
    return Font(name=FONT, **kw)


def cell(ws, ref, value=None, *, bold=False, size=10, color="222222", fill=None, fmt=None, wrap=False,
         align=None, italic=False, border=True):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font(bold=bold, size=size, color=color, italic=italic)
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical="top" if wrap else "center", horizontal=align)
    if border:
        c.border = BOX
    return c


def header(ws, row, values, start_col=1):
    for i, v in enumerate(values):
        ref = f"{get_column_letter(start_col + i)}{row}"
        cell(ws, ref, v, bold=True, color="FFFFFF", fill=HEAD_FILL, wrap=True)


def title(ws, text, sub, width_cols):
    ws["A1"] = text
    ws["A1"].font = font(bold=True, size=15, color=EMERALD)
    ws["A2"] = sub
    ws["A2"].font = font(size=10, color="555555", italic=True)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width_cols)
    ws.row_dimensions[2].height = 30


def q(text: str) -> str:
    """Text as an Excel string literal."""
    return '"' + text.replace('"', '""') + '"'


# ---------------------------------------------------------------------------
# Open-to-buy workbook
# ---------------------------------------------------------------------------

MONTH_COLS = ["D", "E", "F", "G", "H", "I"]
TOTAL_COL = "J"
FIRST_ROW = 13  # first row of the All markets block
BLOCK = 13
LAST_ROW = 400  # SUMIF and COUNTIF ranges reach this far so copied blocks are included


def build_otb(lang: str, t: dict) -> Workbook:
    wb = Workbook()
    sh = t["sheets"]
    guide = wb.active
    guide.title = sh["guide"]
    ws = wb.create_sheet(sh["otb"])
    checks = wb.create_sheet(sh["checks"])
    sizes = wb.create_sheet(sh["sizes"])
    defs = wb.create_sheet(sh["defs"])
    L = t["lines"]
    otb_ref = f"'{sh['otb']}'" if " " in sh["otb"] else sh["otb"]

    # ----- OTB sheet
    title(ws, t["otb_title"], t["otb_sub"], 10)
    for i, (label, val) in enumerate(zip(t["settings"][:4], t["settings_vals"])):
        r = 4 + i
        cell(ws, f"A{r}", label, bold=True)
        ws.merge_cells(f"A{r}:B{r}")
        cell(ws, f"C{r}", val, fill=INPUT_FILL, fmt=PCT if i == 3 else None, color="1F3FBF")
    ws["C6"].comment = Comment(t["cur_note"], t["author"])
    cur, tol = "$C$6", "$C$7"

    # month index and names
    cell(ws, "C10", t["settings"][4], bold=True)
    for i, col in enumerate(MONTH_COLS):
        cell(ws, f"{col}10", i + 1, color="888888", align="center")
    header(ws, 11, t["cols"][:3] + [None] * 6 + [t["cols"][3]])
    for i, col in enumerate(MONTH_COLS):
        cell(ws, f"{col}11", t["months"][i], bold=True, fill=INPUT_FILL, color="1F3FBF", align="center")

    blocks = [(t["total"], t["total_cat"], None)]
    for mi, m in enumerate(t["markets"]):
        for ci, c in enumerate(t["cats"]):
            blocks.append((m, c, EXAMPLE[mi * 2 + ci]))

    rng = lambda col: f"{col}${FIRST_ROW + BLOCK}:{col}${LAST_ROW}"  # detail blocks only
    for bi, (market, cat, ex) in enumerate(blocks):
        top = FIRST_ROW + bi * BLOCK
        is_total = ex is None
        r = {name: top + k for k, name in enumerate(
            ["sales", "md", "close", "open", "onorder", "otb", "margin", "otbcost", "unitcost", "units", "actual", "receipts", "check"])}
        for k in range(BLOCK):
            row = top + k
            fill = TOTAL_FILL if is_total else (BLOCK_FILL if bi % 2 == 0 else None)
            cell(ws, f"A{row}", market if k == 0 else None, bold=k == 0, fill=INPUT_FILL if (k == 0 and not is_total) else fill)
            cell(ws, f"B{row}", cat if k == 0 else None, bold=k == 0, fill=INPUT_FILL if (k == 0 and not is_total) else fill)
            cell(ws, f"C{row}", L[k], bold=k in (5, 12), fill=fill)

        inputs = {"sales", "md", "close", "onorder", "actual", "receipts"}
        for i, col in enumerate(MONTH_COLS):
            prev = MONTH_COLS[i - 1] if i else None
            for name, row in r.items():
                ref = f"{col}{row}"
                if is_total:
                    if name == "margin":
                        f = f"=IF({col}{r['otb']}=0,\"\",1-{col}{r['otbcost']}/{col}{r['otb']})"
                        cell(ws, ref, f, fmt=PCT, fill=TOTAL_FILL)
                    elif name == "unitcost":
                        f = f"=IF({col}{r['units']}=0,\"\",{col}{r['otbcost']}/{col}{r['units']})"
                        cell(ws, ref, f, fmt="#,##0.00", fill=TOTAL_FILL)
                    elif name == "check":
                        f = f"=IF(COUNTIF({col}${FIRST_ROW + BLOCK}:{col}${LAST_ROW},{q(t['flag_over'])})>0,{q(t['flag_over'])},IF(COUNTIF({col}${FIRST_ROW + BLOCK}:{col}${LAST_ROW},{q(t['flag_upd'])})>0,{q(t['flag_upd'])},IF(COUNTIF({col}${FIRST_ROW + BLOCK}:{col}${LAST_ROW},{q(t['flag_drift'])})>0,{q(t['flag_drift'])},IF({col}$10>{cur},\"\",{q(t['ok'])}))))"
                        cell(ws, ref, f, bold=True, fill=TOTAL_FILL, align="center")
                    else:
                        f = f"=SUMIF({rng('$C')},$C{row},{rng(col)})"
                        cell(ws, ref, f, fmt=NUM, fill=TOTAL_FILL, bold=name == "otb")
                    continue
                if name in inputs:
                    if name == "actual":
                        v = ex["actual"][i] if i < len(ex["actual"]) else None
                    elif name == "receipts":
                        v = ex["receipts"][i] if i < len(ex["receipts"]) else None
                    else:
                        v = ex[name][i]
                    cell(ws, ref, v, fmt=NUM, fill=INPUT_FILL, color="1F3FBF")
                elif name == "open":
                    if i == 0:
                        cell(ws, ref, ex["open"], fmt=NUM, fill=INPUT_FILL, color="1F3FBF")
                    else:
                        cell(ws, ref, f"={prev}{r['close']}", fmt=NUM)
                elif name == "otb":
                    f = f"={col}{r['sales']}+{col}{r['md']}+{col}{r['close']}-{col}{r['open']}-{col}{r['onorder']}"
                    cell(ws, ref, f, fmt=NUM, bold=True)
                elif name == "margin":
                    if i == 0:
                        cell(ws, ref, ex["margin"], fmt=PCT, fill=INPUT_FILL, color="1F3FBF")
                    else:
                        cell(ws, ref, f"={prev}{row}", fmt=PCT)
                elif name == "unitcost":
                    if i == 0:
                        cell(ws, ref, ex["cost"], fmt="#,##0.00", fill=INPUT_FILL, color="1F3FBF")
                    else:
                        cell(ws, ref, f"={prev}{row}", fmt="#,##0.00")
                elif name == "otbcost":
                    cell(ws, ref, f"={col}{r['otb']}*(1-{col}{r['margin']})", fmt=NUM)
                elif name == "units":
                    f = f"=IF({col}{r['unitcost']}>0,ROUND({col}{r['otbcost']}/{col}{r['unitcost']},0),0)"
                    cell(ws, ref, f, fmt=NUM)
                elif name == "check":
                    f = (
                        f"=IF({col}{r['otb']}<0,{q(t['flag_over'])},"
                        f"IF({col}$10>{cur},\"\","
                        f"IF(OR({col}{r['actual']}=\"\",{col}{r['receipts']}=\"\"),{q(t['flag_upd'])},"
                        f"IF(AND({col}{r['sales']}>0,ABS({col}{r['actual']}/{col}{r['sales']}-1)>{tol}),{q(t['flag_drift'])},{q(t['ok'])}))))"
                    )
                    cell(ws, ref, f, bold=True, align="center")

        # season totals
        for name, row in r.items():
            ref = f"{TOTAL_COL}{row}"
            fill = TOTAL_FILL if is_total else None
            if name in ("sales", "md", "onorder", "otb", "otbcost", "units", "actual", "receipts"):
                cell(ws, ref, f"=SUM(D{row}:I{row})", fmt=NUM, bold=True, fill=fill)
            elif name == "open":
                cell(ws, ref, f"=D{row}", fmt=NUM, bold=True, fill=fill)
            elif name == "close":
                cell(ws, ref, f"=I{row}", fmt=NUM, bold=True, fill=fill)
            elif name == "margin":
                cell(ws, ref, f"=IF({TOTAL_COL}{r['otb']}=0,\"\",1-{TOTAL_COL}{r['otbcost']}/{TOTAL_COL}{r['otb']})", fmt=PCT, bold=True, fill=fill)
            else:
                cell(ws, ref, None, fill=fill)
        if bi == 1:
            first_units = f"{TOTAL_COL}{r['units']}"

    # conditional formats on every check row
    red = PatternFill("solid", fgColor="F6D5D1")
    amber = PatternFill("solid", fgColor="FBE8C4")
    green = PatternFill("solid", fgColor="D8EDE3")
    area = f"D{FIRST_ROW}:I{LAST_ROW}"
    ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[q(t["flag_over"])], fill=red, font=font(bold=True, color="9B1C1C")))
    ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[q(t["flag_upd"])], fill=amber, font=font(bold=True, color="8A5A00")))
    ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[q(t["flag_drift"])], fill=amber, font=font(bold=True, color="8A5A00")))
    ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[q(t["ok"])], fill=green, font=font(bold=True, color=EMERALD)))
    # negative open-to-buy in red
    ws.conditional_formatting.add(f"D{FIRST_ROW}:J{LAST_ROW}", FormulaRule(formula=[f"AND($C{FIRST_ROW}={q(L[5])},D{FIRST_ROW}<0)"], font=font(bold=True, color="B42318")))

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 30
    for col in MONTH_COLS:
        ws.column_dimensions[col].width = 13
    ws.column_dimensions[TOTAL_COL].width = 15
    ws.freeze_panes = "D12"

    # ----- Checks sheet
    title(checks, t["checks_title"], t["checks_sub"], 6)
    header(checks, 4, t["checks_cols"])
    rows = []
    for bi, (market, cat, ex) in enumerate(blocks[1:], start=1):
        top = FIRST_ROW + bi * BLOCK
        chk_row = top + 12
        rr = 5 + bi - 1
        rows.append(rr)
        cell(checks, f"A{rr}", f"={otb_ref}!A{top}")
        cell(checks, f"B{rr}", f"={otb_ref}!B{top}")
        for j, flag in enumerate([t["flag_over"], t["flag_upd"], t["flag_drift"]]):
            col = "CDE"[j]
            cell(checks, f"{col}{rr}", f"=COUNTIF({otb_ref}!D{chk_row}:I{chk_row},{q(flag)})", fmt="0", align="center")
        cell(checks, f"F{rr}", f"=IF(SUM(C{rr}:E{rr})=0,{q(t['ok'])},{q(t['review'])})", bold=True, align="center")
    tr = rows[-1] + 1
    cell(checks, f"A{tr}", t["total"], bold=True, fill=TOTAL_FILL)
    cell(checks, f"B{tr}", t["total_cat"], bold=True, fill=TOTAL_FILL)
    for col in "CDE":
        cell(checks, f"{col}{tr}", f"=SUM({col}{rows[0]}:{col}{rows[-1]})", fmt="0", bold=True, fill=TOTAL_FILL, align="center")
    cell(checks, f"F{tr}", f"=IF(SUM(C{tr}:E{tr})=0,{q(t['ok'])},{q(t['review'])})", bold=True, fill=TOTAL_FILL, align="center")
    status_area = f"F5:F{tr}"
    checks.conditional_formatting.add(status_area, CellIsRule(operator="equal", formula=[q(t["review"])], fill=amber, font=font(bold=True, color="8A5A00")))
    checks.conditional_formatting.add(status_area, CellIsRule(operator="equal", formula=[q(t["ok"])], fill=green, font=font(bold=True, color=EMERALD)))
    er = tr + 2
    for k, (flag, expl) in enumerate(t["checks_expl"]):
        cell(checks, f"A{er + k}", flag, bold=True)
        cell(checks, f"B{er + k}", expl, wrap=True)
        checks.merge_cells(f"B{er + k}:F{er + k}")
        checks.row_dimensions[er + k].height = 32
    for col, w in zip("ABCDEF", [16, 18, 20, 20, 20, 14]):
        checks.column_dimensions[col].width = w

    # ----- Size and colour split
    title(sizes, t["sizes_title"], t["sizes_sub"], 4)
    cell(sizes, "A4", t["sizes_cat"], bold=True)
    cell(sizes, "B4", f"={otb_ref}!B{FIRST_ROW + BLOCK}", fill=INPUT_FILL)
    cell(sizes, "A5", t["sizes_units"], bold=True)
    cell(sizes, "B5", f"={otb_ref}!{first_units}", fmt=NUM, fill=INPUT_FILL)
    sizes["B5"].comment = Comment(t["sizes_units_note"], t["author"])

    def split_table(start, cols, names, sales):
        header(sizes, start, cols)
        first = start + 1
        last = start + len(names)
        for k, (n, v) in enumerate(zip(names, sales)):
            rr = first + k
            cell(sizes, f"A{rr}", n, fill=INPUT_FILL, color="1F3FBF")
            cell(sizes, f"B{rr}", v, fmt=NUM, fill=INPUT_FILL, color="1F3FBF")
            cell(sizes, f"C{rr}", f"=IF(SUM($B${first}:$B${last})=0,0,B{rr}/SUM($B${first}:$B${last}))", fmt=PCT)
            cell(sizes, f"D{rr}", f"=ROUND($B$5*C{rr},0)", fmt=NUM)
        tr = last + 1
        cell(sizes, f"A{tr}", t["total_row"], bold=True, fill=TOTAL_FILL)
        for col, fmt in zip("BCD", [NUM, PCT, NUM]):
            cell(sizes, f"{col}{tr}", f"=SUM({col}{first}:{col}{last})", fmt=fmt, bold=True, fill=TOTAL_FILL)
        return tr

    end = split_table(7, t["size_cols"], t["sizes"], SIZE_SALES)
    end = split_table(end + 2, t["colour_cols"], t["colours"], COLOUR_SALES)
    sizes[f"A{end + 2}"] = t["rounding_note"]
    sizes[f"A{end + 2}"].font = font(size=9, italic=True, color="666666")
    for col, w in zip("ABCD", [22, 22, 12, 16]):
        sizes.column_dimensions[col].width = w

    # ----- Definitions
    title(defs, t["defs_title"], "", 3)
    header(defs, 4, t["defs_cols"])
    for k, (name, how, owner) in enumerate(t["defs"]):
        rr = 5 + k
        cell(defs, f"A{rr}", name, bold=True, wrap=True)
        cell(defs, f"B{rr}", how, wrap=True)
        cell(defs, f"C{rr}", owner, wrap=True)
        defs.row_dimensions[rr].height = 44
    for col, w in zip("ABC", [28, 70, 20]):
        defs.column_dimensions[col].width = w

    # ----- Guide
    title(guide, t["guide_title"], t["author"], 2)
    rr = 4
    for head, body in t["guide"]:
        if not body:
            cell(guide, f"A{rr}", head, bold=True, color="FFFFFF", fill=HEAD_FILL)
            guide.merge_cells(f"A{rr}:B{rr}")
        else:
            cell(guide, f"A{rr}", head, bold=True, wrap=True)
            cell(guide, f"B{rr}", body, wrap=True)
            guide.row_dimensions[rr].height = 44
        rr += 1
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 90
    guide.sheet_view.showGridLines = False

    for s in (ws, checks, sizes, defs, guide):
        s.sheet_view.zoomScale = 100
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.creator = t["author"]
    wb.properties.title = t["otb_title"]
    return wb


# ---------------------------------------------------------------------------
# Checklist workbook
# ---------------------------------------------------------------------------


def build_checklist(lang: str, t: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = {"en": "Checklist", "pl": "Checklista", "ru": "Чек-лист"}[lang]
    title(ws, t["chk_title"], t["chk_sub"], 7)
    header(ws, 4, t["chk_cols"])
    done, not_yet, na = t["status"]
    dv = DataValidation(type="list", formula1=q(",".join(t["status"])), allow_blank=True)
    ws.add_data_validation(dv)
    first = 5
    for k, (area, check, why) in enumerate(t["chk"]):
        rr = first + k
        cell(ws, f"A{rr}", k + 1, align="center")
        cell(ws, f"B{rr}", area, bold=True, wrap=True)
        cell(ws, f"C{rr}", check, wrap=True)
        cell(ws, f"D{rr}", why, wrap=True, color="555555")
        cell(ws, f"E{rr}", None, fill=INPUT_FILL)
        cell(ws, f"F{rr}", done if k == 0 else not_yet, fill=INPUT_FILL, color="1F3FBF", align="center")
        cell(ws, f"G{rr}", None, fill=INPUT_FILL, wrap=True)
        dv.add(f"F{rr}")
        ws.row_dimensions[rr].height = 46
    last = first + len(t["chk"]) - 1
    s = last + 2
    labels = t["chk_summary"]
    formulas = [
        f"=COUNTIF(F{first}:F{last},{q(done)})",
        f"=COUNT(A{first}:A{last})-COUNTIF(F{first}:F{last},{q(na)})",
        f"=IF(C{s + 1}=0,0,C{s}/C{s + 1})",
        f"=IF(AND(C{s + 1}>0,C{s}=C{s + 1}),{q(t['ready'])},{q(t['not_ready'])})",
    ]
    for k, (lab, f) in enumerate(zip(labels, formulas)):
        cell(ws, f"B{s + k}", lab, bold=True, fill=TOTAL_FILL)
        cell(ws, f"C{s + k}", f, bold=True, fill=TOTAL_FILL, fmt=PCT if k == 2 else "0", align="left")
    green = PatternFill("solid", fgColor="D8EDE3")
    amber = PatternFill("solid", fgColor="FBE8C4")
    grey = PatternFill("solid", fgColor="ECEBE7")
    rng = f"F{first}:F{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[q(done)], fill=green, font=font(bold=True, color=EMERALD)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[q(not_yet)], fill=amber, font=font(bold=True, color="8A5A00")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[q(na)], fill=grey, font=font(color="666666")))
    ws.conditional_formatting.add(f"C{s + 3}", CellIsRule(operator="equal", formula=[q(t["ready"])], fill=green, font=font(bold=True, color=EMERALD)))
    ws.conditional_formatting.add(f"C{s + 3}", CellIsRule(operator="equal", formula=[q(t["not_ready"])], fill=amber, font=font(bold=True, color="8A5A00")))
    for col, w in zip("ABCDEFG", [5, 18, 48, 48, 16, 14, 28]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.creator = t["author"]
    wb.properties.title = t["chk_title"]
    return wb


def main():
    for lang, t in TEXT.items():
        assert len(t["chk"]) == 12, lang
        assert len(t["lines"]) == BLOCK, lang
        folder = os.path.join(OUT, lang)
        os.makedirs(folder, exist_ok=True)
        otb = os.path.join(folder, t["otb_file"])
        chk = os.path.join(folder, t["chk_file"])
        build_otb(lang, t).save(otb)
        build_checklist(lang, t).save(chk)
        print("wrote", os.path.relpath(otb, ROOT), "and", os.path.relpath(chk, ROOT))


if __name__ == "__main__":
    main()
