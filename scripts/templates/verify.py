"""
Evaluates every formula of the built templates with pycel and checks a few
hand-calculated values. Run after build.py:

    python scripts/templates/verify.py
"""
import os
import sys

from openpyxl import load_workbook
from pycel import ExcelCompiler

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build import OUT, TEXT, FIRST_ROW, BLOCK  # noqa: E402

errors = 0


def check(label, got, want):
    global errors
    ok = (abs(got - want) < 1e-6) if isinstance(want, (int, float)) and isinstance(got, (int, float)) else got == want
    if not ok:
        errors += 1
    print(f"  {'ok ' if ok else 'BAD'} {label}: {got!r}" + ("" if ok else f" (want {want!r})"))


def all_formulas(path):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    yield ws.title, c.coordinate


for lang, t in TEXT.items():
    sh = t["sheets"]
    otb_path = os.path.join(OUT, lang, t["otb_file"])
    chk_path = os.path.join(OUT, lang, t["chk_file"])
    print(lang)

    xl = ExcelCompiler(filename=otb_path)
    bad = []
    count = 0
    for sheet, ref in all_formulas(otb_path):
        count += 1
        v = xl.evaluate(f"{sheet}!{ref}")
        if isinstance(v, str) and v.startswith("#"):
            bad.append(f"{sheet}!{ref}={v}")
    print(f"  {count} formulas, {len(bad)} errors {bad[:5]}")
    errors += len(bad)

    o = sh["otb"]
    b1 = FIRST_ROW + BLOCK  # first detail block
    # month 1 of block 1: 42000 + 1500 + 120000 - 110000 - 30000 = 23500
    check("OTB block1 month1", xl.evaluate(f"{o}!D{b1 + 5}"), 23500)
    check("OTB at cost", xl.evaluate(f"{o}!D{b1 + 7}"), 23500 * (1 - 0.62))
    check("OTB units", xl.evaluate(f"{o}!D{b1 + 9}"), round(23500 * 0.38 / 24))
    # month 2 opening stock = month 1 planned closing stock
    check("opening stock month2", xl.evaluate(f"{o}!E{b1 + 3}"), 120000)
    # month 2: 48000 + 1800 + 130000 - 120000 - 40000 = 19800
    check("OTB block1 month2", xl.evaluate(f"{o}!E{b1 + 5}"), 19800)
    check("check block1 month1", xl.evaluate(f"{o}!D{b1 + 12}"), t["ok"])
    check("check block1 month3 (open month)", xl.evaluate(f"{o}!F{b1 + 12}"), "")
    b3 = FIRST_ROW + 3 * BLOCK
    check("check block3 month2 (receipts missing)", xl.evaluate(f"{o}!E{b3 + 12}"), t["flag_upd"])
    b2 = FIRST_ROW + 2 * BLOCK
    # block 2 month 2: actual 21000 vs plan 24000 = -12.5% > 10%
    check("check block2 month2 (drift)", xl.evaluate(f"{o}!E{b2 + 12}"), t["flag_drift"])
    # All markets planned sales month 1 = 42000 + 26000 + 18000 + 11000
    check("total sales month1", xl.evaluate(f"{o}!D{FIRST_ROW}"), 97000)
    check("total OTB season = sum of blocks",
          xl.evaluate(f"{o}!J{FIRST_ROW + 5}"),
          sum(xl.evaluate(f"{o}!J{FIRST_ROW + k * BLOCK + 5}") for k in range(1, 5)))
    c = f"'{sh['checks']}'" if " " in sh["checks"] else sh["checks"]
    check("checks: drift count block2", xl.evaluate(f"{c}!E6"), 1)
    check("checks: not updated block3", xl.evaluate(f"{c}!D7"), 1)
    sz = f"'{sh['sizes']}'" if " " in sh["sizes"] else sh["sizes"]
    check("size shares sum to 1", xl.evaluate(f"{sz}!C13"), 1)

    xl2 = ExcelCompiler(filename=chk_path)
    bad = []
    count = 0
    for sheet, ref in all_formulas(chk_path):
        count += 1
        v = xl2.evaluate(f"{sheet}!{ref}")
        if isinstance(v, str) and v.startswith("#"):
            bad.append(f"{sheet}!{ref}={v}")
    print(f"  checklist: {count} formulas, {len(bad)} errors {bad[:5]}")
    errors += len(bad)
    ws = load_workbook(chk_path).active.title
    check("checklist done", xl2.evaluate(f"{ws}!C18"), 1)
    check("checklist applicable", xl2.evaluate(f"{ws}!C19"), 12)
    check("checklist verdict", xl2.evaluate(f"{ws}!C21"), t["not_ready"])

print("errors:", errors)
sys.exit(1 if errors else 0)
